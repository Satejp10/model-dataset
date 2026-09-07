#!/usr/bin/env python3
"""
add_model.py — insert new model rows into the Models Table workbook.

Handles the mechanical half of adding a model: finding the right sorted
position, matching the surrounding cell styles and number formats, extending
the autofilter range. You supply the researched facts as JSON; this script
places them correctly.

Typical use
-----------
1. Write a spec file (see --example for the shape):
       [{"Model": "GLM-5.3", "Lab": "Z.AI", "Announced": "2026-09", ...}]
2. Preview the insert (writes nothing):
       python logs/add_model.py --spec new.json
3. Apply it:
       python logs/add_model.py --spec new.json --write
4. Then run the normal routine to log and publish the change:
       python logs/diff_dataset.py --write
       python logs/build_constellation.py

Nothing is written unless every model in the spec validates. On any problem the
script prints all of them and exits 1, leaving the workbook untouched.

Only dependency: openpyxl  (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from copy import copy
from pathlib import Path

try:
    import openpyxl
except ModuleNotFoundError:
    sys.exit("openpyxl is required: pip install openpyxl")

LOGS_DIR = Path(__file__).resolve().parent
REPO_ROOT = LOGS_DIR.parent
DEFAULT_XLSX = REPO_ROOT / "Pruned AI Models_Table.xlsx"

SHEET_NAME = "Models"
HEADER_ROW = 2  # the real header row; row 1 holds permalinks/metadata
FIRST_DATA_ROW = HEADER_ROW + 1

REQUIRED = ("Model", "Lab", "Announced")

NUMERIC = {
    "Params (total, B)", "Params (active, B)", "Tokens trained (B)",
    "ALScore", "MMLU", "MMLU-Pro", "GPQA", "HLE", "Count (rough)",
}
BENCHMARKS = {"MMLU", "MMLU-Pro", "GPQA", "HLE"}  # percentages, 0-100
URL_COLUMNS = {"Paper / Repo", "Playground"}

PUBLIC_VALUES = {"🟢", "🔴", "🟡"}
DISCLOSURE_VALUES = {"A", "B", "C", "D", "F"}

DATE_FORMAT = "mmm/yyyy"

EXAMPLE = [
    {
        "Model": "GLM-5.3",
        "Lab": "Z.AI",
        "Params (total, B)": 355,
        "Params (active, B)": 32,
        "Announced": "2026-09",
        "Arch": "MoE",
        "Tokens trained (B)": 23000,
        "ALScore": 41.2,
        "MMLU-Pro": 86.1,
        "GPQA": 79.4,
        "Training dataset": "synthetic, web-scale",
        "Public?": "🟢",
        "Disclosure score": "B",
        "Paper / Repo": "https://example.org/glm-5-3",
        "Tags": "Reasoning",
        "Notes": "One or two sentences on what it is and what shipped.",
        "Playground": "https://example.org/chat",
    }
]


# --------------------------------------------------------------------------- #
# Sheet plumbing
# --------------------------------------------------------------------------- #
def clean_header(name) -> str:
    """Tidy a header cell: drop sort glyphs, collapse embedded newlines/spaces."""
    s = str(name).replace("​", "")
    s = re.sub(r"[▼▲△▽]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return re.sub(r"(\w) -(\w)", r"\1-\2", s)  # rejoin a hyphenated word split across lines


def read_columns(ws) -> dict[str, int]:
    """Map cleaned column name -> 1-based column index, skipping the blank spacer."""
    columns = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(row=HEADER_ROW, column=c).value
        if raw is None:
            continue
        name = clean_header(raw)
        if name:
            columns[name] = c
    return columns


def last_data_row(ws, model_col: int) -> int:
    """Last row holding a model name."""
    last = HEADER_ROW
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(row=r, column=model_col).value not in (None, ""):
            last = r
    return last


def existing_models(ws, columns: dict[str, int], last: int) -> set[tuple[str, str]]:
    model_col, lab_col = columns["Model"], columns["Lab"]
    found = set()
    for r in range(FIRST_DATA_ROW, last + 1):
        name = ws.cell(row=r, column=model_col).value
        if name in (None, ""):
            continue
        lab = ws.cell(row=r, column=lab_col).value or ""
        found.add((str(name).strip().casefold(), str(lab).strip().casefold()))
    return found


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def parse_announced(value) -> _dt.datetime:
    """Accept YYYY-MM or YYYY-MM-DD; the sheet stores every model at day 1."""
    if isinstance(value, _dt.datetime):
        return value.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if isinstance(value, _dt.date):
        return _dt.datetime(value.year, value.month, 1)
    text = str(value).strip()
    for pattern in ("%Y-%m", "%Y-%m-%d"):
        try:
            return _dt.datetime.strptime(text, pattern).replace(day=1)
        except ValueError:
            continue
    raise ValueError(f"expected YYYY-MM or YYYY-MM-DD, got {value!r}")


def validate(spec: dict, index: int, columns: dict[str, int], seen: set) -> tuple[dict, list[str]]:
    """Return (normalised record, errors). Record is only usable when errors is empty."""
    errors: list[str] = []
    where = f"spec[{index}]"

    if not isinstance(spec, dict):
        return {}, [f"{where}: expected an object, got {type(spec).__name__}"]

    label = str(spec.get("Model", "?")).strip() or "?"
    where = f"spec[{index}] ({label})"

    unknown = [k for k in spec if k not in columns]
    if unknown:
        errors.append(f"{where}: unknown column(s) {unknown}. Valid: {sorted(columns)}")

    for field in REQUIRED:
        if str(spec.get(field, "")).strip() == "":
            errors.append(f"{where}: '{field}' is required")

    record: dict[str, object] = {}
    for key, raw in spec.items():
        if key not in columns:
            continue  # already reported
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            continue  # blank stays blank

        if key == "Announced":
            try:
                record[key] = parse_announced(raw)
            except ValueError as exc:
                errors.append(f"{where}: 'Announced' {exc}")
        elif key in NUMERIC:
            try:
                number = float(str(raw).replace(",", "").strip())
            except ValueError:
                errors.append(f"{where}: '{key}' must be a number, got {raw!r}")
                continue
            if number <= 0:
                errors.append(f"{where}: '{key}' must be positive, got {raw!r}")
            elif key in BENCHMARKS and number > 100:
                errors.append(f"{where}: '{key}' is a percentage but got {raw!r} "
                              f"(a misplaced decimal point?)")
            else:
                record[key] = int(number) if number.is_integer() else number
        else:
            text = str(raw).strip()
            if key == "Public?" and text not in PUBLIC_VALUES:
                errors.append(f"{where}: 'Public?' must be one of {sorted(PUBLIC_VALUES)}, got {text!r}")
            elif key == "Disclosure score" and text.upper() not in DISCLOSURE_VALUES:
                errors.append(f"{where}: 'Disclosure score' must be one of "
                              f"{sorted(DISCLOSURE_VALUES)}, got {text!r}")
            elif key in URL_COLUMNS and not text.startswith(("http://", "https://")):
                errors.append(f"{where}: '{key}' must be a URL, got {text!r}")
            else:
                record[key] = text.upper() if key == "Disclosure score" else text

    if not errors:
        identity = (record["Model"].casefold(), record["Lab"].casefold())
        if identity in seen:
            errors.append(f"{where}: '{record['Model']}' by {record['Lab']} is already in the sheet "
                          f"(pass --allow-duplicate to add it anyway)")
        else:
            seen.add(identity)

    return record, errors


# --------------------------------------------------------------------------- #
# Insertion
# --------------------------------------------------------------------------- #
def find_position(ws, announced_col: int, announced: _dt.datetime, last: int) -> int:
    """Row index for a new model: top of its month's block, keeping newest-first order."""
    for r in range(FIRST_DATA_ROW, last + 1):
        value = ws.cell(row=r, column=announced_col).value
        if not isinstance(value, (_dt.datetime, _dt.date)):
            continue
        current = value if isinstance(value, _dt.datetime) else _dt.datetime(value.year, value.month, value.day)
        if (current.year, current.month) <= (announced.year, announced.month):
            return r
    return last + 1


def insert_record(ws, columns: dict[str, int], record: dict, last: int, max_col: int) -> int:
    """Insert one record, inheriting the styles of the row it displaces."""
    row = find_position(ws, columns["Announced"], record["Announced"], last)
    anchor = row if row <= last else last
    # insert_rows() leaves the new row unstyled, and _style is the only handle that
    # carries font, fill, border, alignment and number format across in one piece
    styles = [copy(ws.cell(row=anchor, column=c)._style) for c in range(1, max_col + 1)]

    ws.insert_rows(row)
    for c in range(1, max_col + 1):
        ws.cell(row=row, column=c)._style = styles[c - 1]

    for name, value in record.items():
        ws.cell(row=row, column=columns[name]).value = value
    ws.cell(row=row, column=columns["Announced"]).number_format = DATE_FORMAT
    return row


def extend_autofilter(ws, last: int) -> str | None:
    """Grow the autofilter to cover the new final row."""
    ref = ws.auto_filter.ref
    if not ref:
        return None
    match = re.fullmatch(r"(\$?[A-Z]+\$?\d+):(\$?[A-Z]+)\$?(\d+)", ref)
    if not match:
        return None
    start, end_col, _ = match.groups()
    ws.auto_filter.ref = f"{start}:{end_col}${last}"
    return ws.auto_filter.ref


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def load_spec(source: str) -> list:
    if source == "-":
        text = sys.stdin.read()
    else:
        path = Path(source)
        if not path.exists():
            sys.exit(f"Spec file not found: {path}")
        text = path.read_text(encoding="utf-8")
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        sys.exit(f"Spec is not valid JSON: {exc}")
    return data if isinstance(data, list) else [data]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", help="JSON file describing the model(s) to add, or - for stdin")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX,
                    help=f"workbook to edit (default: {DEFAULT_XLSX.name})")
    ap.add_argument("--write", action="store_true",
                    help="save the workbook (default: preview only)")
    ap.add_argument("--allow-duplicate", action="store_true",
                    help="permit a model whose name and lab already appear in the sheet")
    ap.add_argument("--example", action="store_true",
                    help="print a spec template and exit")
    args = ap.parse_args()

    if args.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return
    if not args.spec:
        sys.exit("Nothing to do: pass --spec PATH (or --example to see the shape).")
    if not args.xlsx.exists():
        sys.exit(f"Workbook not found: {args.xlsx}")

    specs = load_spec(args.spec)
    if not specs:
        sys.exit("Spec contains no models.")

    wb = openpyxl.load_workbook(args.xlsx)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"{args.xlsx.name}: no sheet named '{SHEET_NAME}'")
    ws = wb[SHEET_NAME]

    columns = read_columns(ws)
    for field in REQUIRED:
        if field not in columns:
            sys.exit(f"{args.xlsx.name}: header row {HEADER_ROW} has no '{field}' column")

    last = last_data_row(ws, columns["Model"])
    seen = set() if args.allow_duplicate else existing_models(ws, columns, last)

    records, errors = [], []
    for i, spec in enumerate(specs):
        record, problems = validate(spec, i, columns, seen)
        errors.extend(problems)
        if not problems:
            records.append(record)

    if errors:
        print(f"{len(errors)} problem(s) found — nothing was written:\n", file=sys.stderr)
        for message in errors:
            print(f"  ! {message}", file=sys.stderr)
        sys.exit(1)

    # Insert oldest-listed first so that models sharing a month keep spec order.
    max_col = ws.max_column
    placements = []
    for record in reversed(records):
        row = insert_record(ws, columns, record, last, max_col)
        last += 1
        # earlier inserts at or below this row have just been pushed down one
        placements = [(r + 1 if r >= row else r, rec) for r, rec in placements]
        placements.append((row, record))
    placements.reverse()

    filter_ref = extend_autofilter(ws, last)

    print(f"{len(records)} model(s) to add · sheet grows to {last - HEADER_ROW} rows")
    if filter_ref:
        print(f"autofilter → {filter_ref}")
    print()
    for row, record in placements:
        filled = [k for k in record if k not in REQUIRED]
        print(f"  + row {row}: {record['Model']} — {record['Lab']} "
              f"({record['Announced']:%Y-%m}) · {len(filled)} extra field(s)")
        blank = [c for c in columns if c not in record]
        if blank:
            print(f"      left blank: {', '.join(sorted(blank))}")

    if not args.write:
        print("\nPreview only. Re-run with --write to save the workbook.")
        return

    wb.save(args.xlsx)
    print(f"\nSaved {args.xlsx.name}. Next: python logs/diff_dataset.py --write")


if __name__ == "__main__":
    main()
